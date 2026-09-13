"""Таблицы выгрузки FineReader: чтение DOCX в порядке документа.

ПОЧЕМУ НЕ ``Document.tables``. Он отдаёт таблицы списком, оторванным от текста вокруг, а
текст вокруг здесь нужен: у половины испорченных таблиц собственные ячейки — сплошной мусор,
и привязать такую таблицу к странице можно только по соседним абзацам («Из приведенных
данных видно...»). Поэтому обход идёт по телу документа, как в ``ocr_utils.docx_md``.

ОБЪЕДИНЁННЫЕ ЯЧЕЙКИ. python-docx в ``row.cells`` повторяет объединённую ячейку столько раз,
сколько колонок она накрывает. Здесь такие повторы схлопываются, а факт объединения
запоминается: в многоуровневых шапках («Хвойные породы» над двумя графами) боковой текст
живёт как раз под объединёнными заголовками.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

# Сколько абзацев по обе стороны таблицы запоминать. Два: первый обычно вводит таблицу
# («что видно из табл. 2»), второй даёт достаточно редких слов для поиска страницы.
CONTEXT_PARAGRAPHS = 2


@dataclass
class DocxTable:
    """Одна таблица выгрузки вместе с окрестностями."""

    issue: str
    index: int  # порядковый номер таблицы в документе, с нуля
    rows: list[list[str]]  # текст ячеек; объединённые не повторяются
    merged: set[tuple[int, int]] = field(default_factory=set)  # клетки, которые были повторами
    context_before: list[str] = field(default_factory=list)
    context_after: list[str] = field(default_factory=list)
    docx_page_hint: int = 0  # счётчик страниц DOCX: только окно поиска, не ответ

    @property
    def n_rows(self) -> int:
        return len(self.rows)

    @property
    def n_cols(self) -> int:
        return max((len(row) for row in self.rows), default=0)

    def cells(self) -> list[str]:
        return [text for row in self.rows for text in row]

    def head_cells(self, rows: int = 2) -> list[str]:
        return [text for row in self.rows[:rows] for text in row]


def _cell_texts(row) -> tuple[list[str], set[int]]:
    """Тексты ячеек строки без повторов объединённых; вторым — индексы схлопнутых."""
    texts: list[str] = []
    merged: set[int] = set()
    seen: set[int] = set()
    for cell in row.cells:
        marker = id(cell._tc)
        if marker in seen:
            merged.add(len(texts) - 1 if texts else 0)
            continue
        seen.add(marker)
        texts.append(cell.text.strip())
    return texts, merged


def _paragraph_text(element) -> str:
    return "".join(node.text or "" for node in element.iter(qn("w:t"))).strip()


def _page_breaks(element) -> int:
    """Сколько разрывов страниц внутри элемента.

    Счётчик по разрывам и несплошным секциям — то, что в этих файлах вообще есть; но он
    расходится с числом страниц PDF на ±1 (см. ``pages``), поэтому используется как окно.
    """
    breaks = sum(1 for node in element.findall(".//" + qn("w:br")) if node.get(qn("w:type")) == "page")
    section = element.find(".//" + qn("w:sectPr"))
    if section is not None:
        kind = section.find(qn("w:type"))
        if kind is None or kind.get(qn("w:val")) != "continuous":
            breaks += 1
    return breaks


def iter_tables(docx_path: Path, issue: str) -> list[DocxTable]:
    """Все таблицы документа в порядке чтения, с окрестностями и подсказкой о странице."""
    document = Document(str(docx_path))
    body = document.element.body
    tables_xml = {id(table._tbl): table for table in document.tables}

    found: list[DocxTable] = []
    recent: list[str] = []
    page = 1
    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            text = _paragraph_text(child)
            if text:
                recent.append(text)
                recent[:] = recent[-CONTEXT_PARAGRAPHS:]
            page += _page_breaks(child)
            if found and len(found[-1].context_after) < CONTEXT_PARAGRAPHS and text:
                found[-1].context_after.append(text)
        elif tag == "tbl":
            table = tables_xml.get(id(child))
            if table is None:
                continue
            rows: list[list[str]] = []
            merged: set[tuple[int, int]] = set()
            for row_index, row in enumerate(table.rows):
                texts, merged_columns = _cell_texts(row)
                rows.append(texts)
                merged.update((row_index, column) for column in merged_columns)
            found.append(
                DocxTable(
                    issue=issue,
                    index=len(found),
                    rows=rows,
                    merged=merged,
                    context_before=list(recent),
                    docx_page_hint=page,
                )
            )
            recent = []
            page += _page_breaks(child)
    return found
